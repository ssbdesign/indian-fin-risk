
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, timedelta
file_path = Path(__file__).resolve().parent.parent / "data" / "raw" / "rbi_repo_rate" / "major_monetary_policy_rates.xlsx"

if not file_path.exists():
    raise FileNotFoundError(f"Workbook not found: {file_path.resolve()}")

workbook = load_workbook(file_path, data_only=True, read_only=True)

print("Workbook sheets:", workbook.sheetnames)

for sheet in workbook.worksheets:
    print(f"\n--- Sheet: {sheet.title} ---")

    for row in sheet.iter_rows(min_row=1, max_row=15, values_only=True):
        values = list(row[:8])

        # Convert Excel serial dates in the displayed output only.
        for i, value in enumerate(values):
            if isinstance(value, (int, float)) and i == 1:
                try:
                    values[i] = (
                        datetime(1899, 12, 30) + timedelta(days=value)
                    ).strftime("%Y-%m-%d")
                except (OverflowError, ValueError):
                    pass

        print(values)

workbook.close()